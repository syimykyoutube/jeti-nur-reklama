"""v1.1 UI contracts, safe HTML, fallback deletion and explicit DOCX layout."""

from datetime import date
from zoneinfo import ZoneInfo
from html import escape
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import aiosqlite
from aiogram.exceptions import TelegramBadRequest
from aiogram.methods import DeleteMessage, EditMessageText, SetMyCommands
from aiogram.types import Animation, Audio, Contact, Document as TelegramDocument, Location, PhotoSize, Sticker, Video, VideoNote, Voice
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from lxml import etree
from PIL import Image

from bot.callbacks import AnalyticsPeriodCallback, DeleteOrderCallback, OrderConfirmCallback, OrderPageCallback, PeriodChoiceCallback, ProductCallback
from bot.handlers.analytics import AnalyticsStates
from bot.handlers.catalog import CatalogStates
from bot.handlers.orders import OrderStates
from bot.i18n import TEXTS, t
from bot.main import create_dispatcher, register_commands
from bot.services.charts import chart_png, plt
from bot.services.report_docx import create_report
from bot.services.stats import ChartBucket, ChartSeries, make_period, validate_days
from test_stage1 import BotTestCase, NOW

PHOTO = [PhotoSize(file_id="photo", file_unique_id="unique", width=10, height=10)]
RECEIPT = TelegramDocument(file_id="file", file_unique_id="unique", file_name="receipt.pdf")


class UIv11Tests(BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        await self.database.set_user_language(101, "Aibek", "ru")
        await self.database.set_user_language(202, "Aibek", "ky")
        self.product_id = await self.database.add_product("Машина", 200)

    def state(self, user_id=101):
        return self.dispatcher.fsm.get_context(bot=self.bot, chat_id=user_id, user_id=user_id)

    def deleted(self):
        return [call for call in self.session.calls if isinstance(call, DeleteMessage)]

    async def order(self, name="Машина", author="Aibek"):
        with patch("bot.db.datetime") as clock:
            clock.now.return_value = NOW
            return await self.database.add_order(product_id=self.product_id, product_name=name, price=200,
                created_by=101, created_by_name=author, receipt_file_id="file", receipt_file_unique_id="unique", receipt_type="photo")

    async def test_main_keyboard_two_vertical_buttons_and_no_repeated_headers(self):
        await self.send_message("/menu")
        menu = self.sent_messages()[-1]
        self.assertEqual([[button.text for button in row] for row in menu.reply_markup.keyboard], [["📋 Каталог"], ["📊 Аналитика"]])
        self.assertTrue(menu.reply_markup.resize_keyboard)
        self.assertTrue(menu.reply_markup.is_persistent)
        for command in ("/language", t("menu_catalog", "ru"), t("menu_analytics", "ru"), "/deleteorder"):
            self.session.calls.clear()
            await self.send_message(command)
            self.assertFalse(any(message.text in [t("main_menu", "ru"), "📋 Каталог", "📊 Аналитика", t("menu_delete_order", "ru")]
                                 for message in self.sent_messages()))
            for message in self.sent_messages():
                if message.reply_markup and hasattr(message.reply_markup, "inline_keyboard"):
                    self.assertFalse(any(button.text == "Меню" for row in message.reply_markup.inline_keyboard for button in row))

    async def test_command_list_has_seven_commands_and_localized_edit_deleteorder(self):
        await register_commands(self.bot, "ky")
        await register_commands(self.bot, "ru", "ru")
        commands = [call for call in self.session.calls if isinstance(call, SetMyCommands)]
        self.assertEqual([command.command for command in commands[0].commands], ["start", "menu", "language", "edit", "deleteorder", "myid", "cancel"])
        self.assertIsNone(commands[0].language_code)
        self.assertEqual(commands[1].language_code, "ru")
        self.assertEqual(commands[0].commands[4].description, "Заказды өчүрүү")
        self.assertEqual(commands[1].commands[4].description, "Удалить заказ")

    async def test_unexpected_types_and_idle_text_photo_document_are_deleted(self):
        files = {"file_id": "file", "file_unique_id": "unique"}
        content = [
            ("hello", {}), (None, {"photo": PHOTO}), (None, {"document": RECEIPT}),
            (None, {"voice": Voice(**files, duration=1)}),
            (None, {"video": Video(**files, width=10, height=10, duration=1)}),
            (None, {"video_note": VideoNote(**files, length=10, duration=1)}),
            (None, {"audio": Audio(**files, duration=1)}),
            (None, {"animation": Animation(**files, width=10, height=10, duration=1), "document": RECEIPT}),
            (None, {"sticker": Sticker(**files, type="regular", width=10, height=10, is_animated=False, is_video=False)}),
            (None, {"contact": Contact(phone_number="123", first_name="Test")}),
            (None, {"location": Location(latitude=42, longitude=74)}),
        ]
        for text, data in content:
            await self.send_message(text, **data)
        self.assertEqual(len(self.deleted()), len(content))
        self.assertEqual(self.sent_messages(), [])
        self.assertEqual(self.dispatcher.sub_routers[-1].name, "fallback")

    async def test_delete_failures_are_suppressed(self):
        original = self.session.make_request
        async def unavailable(bot, method, timeout=None):
            if isinstance(method, DeleteMessage):
                raise TelegramBadRequest(method=method, message="message can't be deleted")
            return await original(bot, method, timeout)
        with patch.object(self.session, "make_request", side_effect=unavailable):
            await self.send_message("unwanted")
        await self.send_message("/menu")
        self.assertEqual(self.sent_messages()[-1].text, t("main_menu", "ru"))

    async def test_expected_fsm_text_and_receipts_are_not_deleted(self):
        await self.click("catalog:add")
        await self.send_message("Card <new>")
        await self.send_message("200")
        self.assertEqual(self.deleted(), [])
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        await self.send_message(None, photo=PHOTO)
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        await self.send_message(None, document=RECEIPT)
        self.assertEqual(self.deleted(), [])
        self.assertEqual((await self.database.get_orders_page()).total, 2)

    async def test_wrong_type_during_fsm_deletes_without_advancing(self):
        await self.click("catalog:add")
        await self.send_message(None, photo=PHOTO)
        self.assertEqual(await self.state().get_state(), CatalogStates.add_name.state)
        self.assertEqual(len(self.deleted()), 1)
        await self.click(PeriodChoiceCallback(period="custom").pack())
        await self.send_message(None, document=RECEIPT)
        self.assertEqual(await self.state().get_state(), AnalyticsStates.custom_days.state)
        self.assertEqual(len(self.deleted()), 2)
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        draft = await self.state().get_data()
        await self.send_message("wrong receipt")
        self.assertEqual(len(self.deleted()), 3)
        self.assertEqual(await self.state().get_data(), draft)
        self.assertEqual(self.sent_messages()[-1].text, t("receipt_wrong", "ru"))

    async def test_known_commands_and_menu_work_during_custom_days(self):
        for command in ("/start", "/menu", "/language", "/edit", "/deleteorder", "/myid", "/cancel", "/backup", "📋 Каталог", "📊 Аналитика"):
            await self.click(PeriodChoiceCallback(period="custom").pack())
            self.session.calls.clear()
            await self.send_message(command)
            self.assertEqual(self.deleted(), [])
            self.assertTrue(self.sent_messages())
            if command not in {"/myid", "/backup"}:
                self.assertNotEqual(await self.state().get_state(), AnalyticsStates.custom_days.state)

    async def test_custom_days_retries_invalid_accepts_bounds_and_cancel(self):
        for user_id, lang in ((101, "ru"), (202, "ky")):
            await self.click(PeriodChoiceCallback(period="custom").pack(), user_id)
            for text in ("0", "100", "-1", "2.5", "1,5", "abc", " ", "١"):
                await self.send_message(text, user_id)
                self.assertEqual(await self.state(user_id).get_state(), AnalyticsStates.custom_days.state)
                self.assertEqual(self.sent_messages()[-1].text, t("ask_days", lang))
            for days in ("1", "99"):
                await self.send_message(days, user_id)
                self.assertIsNone(await self.state(user_id).get_state())
                keyboard = self.sent_messages()[-1].reply_markup.inline_keyboard
                data = AnalyticsPeriodCallback.unpack(keyboard[0][0].callback_data)
                self.assertEqual(data.period, "t" if days == "1" else "d99")
                await self.click(PeriodChoiceCallback(period="custom").pack(), user_id)
            await self.send_message("/cancel", user_id)
            self.assertIsNone(await self.state(user_id).get_state())
            self.assertEqual(self.sent_messages()[-1].text, t("cancelled", lang))
        self.assertEqual(self.deleted(), [])

    async def test_callbacks_survive_restart_and_share_selected_bounds(self):
        await self.order()
        with patch("bot.handlers.analytics.make_period", side_effect=lambda key, tz: make_period(key, tz, NOW)):
            await self.click(PeriodChoiceCallback(period="d7").pack())
        rows = self.sent_messages()[-1].reply_markup.inline_keyboard
        actions = [button.callback_data for row in rows[:3] for button in row]
        for data in actions:
            payload = AnalyticsPeriodCallback.unpack(data)
            self.assertEqual((payload.period, payload.anchor), ("d7", make_period("7", self.config.timezone, NOW).anchor))
            self.assertLessEqual(len(data.encode()), 64)
        await self.dispatcher.fsm.close()
        self.dispatcher = create_dispatcher(self.config, self.database)
        for data in actions:
            self.session.calls.clear()
            await self.click(data)
            self.assertFalse(any(message.text == t("choose_period", "ru") for message in self.sent_messages()))
            self.assertIsNone(await self.state().get_state())
        self.assertEqual(len(self.session.uploaded_documents), 1)

    async def test_html_escapes_product_names_and_authors_in_all_views(self):
        name = '<b>A & B</b>'
        author = '<i>Name & surname</i>'
        await self.database.update_product_name(self.product_id, name)
        await self.send_message("📋 Каталог")
        self.assertEqual(self.sent_messages()[0].text, f"<b>{escape(name)} — 200 сом</b>")
        self.assertEqual(self.bot.default.parse_mode, "HTML")
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        self.assertIn(escape(name), self.sent_messages()[-1].text)
        await self.send_message("/cancel")
        await self.order(name, author)
        for data in (AnalyticsPeriodCallback(mode="top", period="all").pack(), AnalyticsPeriodCallback(mode="orders", period="all").pack()):
            self.session.calls.clear()
            await self.click(data)
            self.assertIn(escape(name), self.sent_messages()[0].text)
        self.assertIn(escape(author), self.sent_messages()[0].text)
        self.session.calls.clear()
        await self.send_message("/deleteorder")
        self.assertIn(escape(name), self.sent_messages()[0].text)
        self.assertIn(escape(author), self.sent_messages()[0].text)

    async def test_delete_edits_same_card_no_cancels_yes_soft_deletes(self):
        order_id = await self.order()
        await self.send_message("/deleteorder")
        data = DeleteOrderCallback(order_id=order_id, page=0).pack()
        self.session.calls.clear()
        await self.click(data, message_id=777)
        token = (await self.state().get_data())["token"]
        edit = self.session.calls[-1]
        self.assertIsInstance(edit, EditMessageText)
        self.assertEqual((edit.message_id, edit.text), (777, t("confirm_delete", "ru")))
        await self.click(OrderConfirmCallback(action="yes", token=token).pack(), message_id=888)
        self.assertIsNotNone(await self.database.get_active_order(order_id))
        await self.click(OrderConfirmCallback(action="no", token=token).pack(), message_id=777)
        self.assertEqual(self.session.calls[-1].text, t("delete_cancelled", "ru"))
        self.assertIsNone(self.session.calls[-1].reply_markup)
        self.assertEqual(self.session.calls[-1].message_id, 777)
        await self.click(data, message_id=777)
        token = (await self.state().get_data())["token"]
        await self.click(OrderConfirmCallback(action="yes", token=token).pack(), message_id=777)
        self.assertEqual(self.session.calls[-1].text, t("order_deleted", "ru"))
        self.assertEqual(self.session.calls[-1].message_id, 777)
        self.assertEqual(self.sent_messages(), [])
        self.assertIsNone(await self.database.get_active_order(order_id))

    async def test_numbering_and_arrows_only_above_ten_in_both_lists(self):
        ids = [await self.order() for _ in range(12)]
        for deletion in (True, False):
            self.session.calls.clear()
            if deletion:
                await self.send_message("/deleteorder")
            else:
                await self.click(AnalyticsPeriodCallback(mode="orders", period="all").pack())
            messages = self.sent_messages()
            self.assertEqual(len(messages), 11)
            self.assertTrue(messages[0].text.startswith("<b>1."))
            self.assertTrue(messages[9].text.startswith("<b>10."))
            next_page = messages[-1].reply_markup.inline_keyboard[0][0].callback_data
            self.session.calls.clear()
            await self.click(next_page)
            messages = self.sent_messages()
            self.assertEqual(len(messages), 3)
            self.assertTrue(messages[0].text.startswith("<b>11."))
            self.assertTrue(messages[1].text.startswith("<b>12."))
            self.assertEqual(messages[-1].reply_markup.inline_keyboard[0][0].text, "◀")
        for order_id in ids[:2]:
            await self.database.soft_delete_order(order_id, 101)
        for deletion in (True, False):
            self.session.calls.clear()
            if deletion:
                await self.send_message("/deleteorder")
            else:
                await self.click(AnalyticsPeriodCallback(mode="orders", period="all").pack())
            self.assertEqual(len(self.sent_messages()), 10)
            self.assertFalse(any(message.text == t("pagination", "ru") for message in self.sent_messages()))

    async def test_empty_views_keep_change_period_button(self):
        for mode in ("summary", "top", "chart", "orders"):
            await self.click(AnalyticsPeriodCallback(mode=mode, period="all").pack(), 202)
            result = self.sent_messages()[-1]
            self.assertIn("Бул убакытта заказ жок", result.text)
            self.assertEqual(result.reply_markup.inline_keyboard[-1][0].text, "🔄 Убакытты өзгөртүү")

    async def test_docx_media_references_colors_widths_and_footer(self):
        for _ in range(2):
            await self.order()
        path = Path(self.temp.name) / "report.docx"
        await create_report(self.database, path, make_period("all", self.config.timezone, NOW), self.config.timezone, "ky", NOW)
        document = Document(path)
        self.assertAlmostEqual(document.sections[0].page_width / Cm(1), 21, places=2)
        self.assertAlmostEqual(document.sections[0].left_margin / Cm(1), 2, places=2)
        self.assertEqual(document.paragraphs[0].runs[0].font.size.pt, 22)
        self.assertEqual(document.paragraphs[0].runs[0].font.color.rgb.__str__(), "1F4E79")
        self.assertNotIn("Asia/Bishkek", "\n".join(p.text for p in document.paragraphs))
        self.assertEqual(len(document.tables), 3)
        self.assertEqual([round(column.width / Cm(1), 1) for column in document.tables[1].columns], [1.2, 9.8, 2.5, 3.5])
        self.assertEqual([round(column.width / Cm(1), 1) for column in document.tables[2].columns], [1.2, 4.5, 4.8, 3.0, 3.5])
        self.assertEqual(document.tables[2].rows[-1].cells[0].text, "Жалпы: 400 сом")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        with ZipFile(path) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
            self.assertEqual(xml.find("w:background", ns).get(qn("w:color")), "FFFFFF")
            settings = etree.fromstring(archive.read("word/settings.xml"))
            self.assertIsNotNone(settings.find("w:displayBackgroundShape", ns))
            for run in xml.findall(".//w:r", ns):
                color = run.find("w:rPr/w:color", ns)
                self.assertIsNotNone(color)
                self.assertIn(color.get(qn("w:val")), {"1F1F1F", "FFFFFF", "1F4E79"})
            for cell in xml.findall(".//w:tc", ns):
                self.assertIsNotNone(cell.find("w:tcPr/w:shd", ns))
            media = [member for member in archive.namelist() if member.startswith("word/media/") and member.endswith(".png")]
            self.assertEqual(len(media), 1)
            with Image.open(BytesIO(archive.read(media[0]))) as image:
                image.load()
                self.assertEqual(image.size, (1800, 1000))
                self.assertAlmostEqual(image.info["dpi"][0], 200, delta=1)
                self.assertEqual(image.convert("RGB").getpixel((0, 0)), (255, 255, 255))
                self.assertGreater(len(image.convert("RGB").getcolors(image.width * image.height)), 100)
            image_id = xml.find(".//a:blip", ns).get(qn("r:embed"))
            relations = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
            relation = next(item for item in relations if item.get("Id") == image_id)
            self.assertEqual("word/" + relation.get("Target"), media[0])
            self.assertEqual(document.inline_shapes[0].width, Cm(16))
            self.assertIn(" PAGE ", archive.read("word/footer1.xml").decode())
            self.assertIn('w:val="center"', archive.read("word/footer1.xml").decode())

    async def test_inline_cancel_keeps_menu_and_only_cancels_catalog_forms(self):
        await self.click("catalog:add")
        button = self.sent_messages()[-1].reply_markup.inline_keyboard[0][0]
        self.assertEqual(button.callback_data, "catalog:cancel")
        await self.send_message("Draft")
        await self.click(button.callback_data)
        self.assertIsNone(await self.state().get_state())
        self.assertEqual(len(await self.database.list_active_products()), 1)
        await self.click(ProductCallback(action="edit", product_id=self.product_id).pack())
        from bot.callbacks import CatalogFormCallback
        token = (await self.state().get_data())["token"]
        await self.click(CatalogFormCallback(action="price", token=token).pack())
        await self.click("catalog:cancel")
        self.assertIsNone(await self.state().get_state())
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        await self.click("catalog:cancel")
        self.assertEqual(await self.state().get_state(), OrderStates.awaiting_receipt.state)
        await self.click("order:cancel")
        self.assertIsNone(await self.state().get_state())
        self.assertEqual((await self.database.get_orders_page()).total, 0)

    async def test_no_cancels_original_confirmation_when_new_orders_arrive(self):
        first = await self.order()
        await self.send_message("/deleteorder")
        original = self.sent_messages()[-1]
        data = original.reply_markup.inline_keyboard[0][0].callback_data
        self.assertEqual(DeleteOrderCallback.unpack(data).number, 1)
        await self.order()
        await self.click(data)
        token = (await self.state().get_data())["token"]
        await self.click(OrderConfirmCallback(action="no", token=token).pack())
        self.assertEqual(self.session.calls[-1].text, t("delete_cancelled", "ru"))
        self.assertIsNone(self.session.calls[-1].reply_markup)
        self.assertIsNotNone(await self.database.get_active_order(first))

    async def test_schema_stays_unchanged_after_ui_operations(self):
        async def schema():
            async with aiosqlite.connect(self.path) as connection:
                async with connection.execute("SELECT type, name, sql FROM sqlite_master ORDER BY name") as cursor:
                    return await cursor.fetchall()
        before = await schema()
        await self.send_message("/menu")
        await self.click(PeriodChoiceCallback(period="custom").pack())
        await self.send_message("15")
        await self.send_message("/deleteorder")
        self.assertEqual(await schema(), before)


class v11ValidationTests(unittest.TestCase):
    def test_days_1_to_99_and_invalid_callback_tokens(self):
        for days in range(1, 100):
            self.assertEqual(validate_days(str(days)), str(days))
            self.assertEqual(make_period(str(days), ZoneInfo("Asia/Bishkek"), NOW).key, str(days))
        for text in (None, "0", "100", "-5", "NaN", "4.0", "4e1", "١", "+4"):
            with self.assertRaises(ValueError):
                validate_days(text)
        for token in ("d0", "d100", "d01", "365", "invalid", ""):
            with self.assertRaises(ValueError):
                AnalyticsPeriodCallback.unpack(f"an:top:{token}:0:0")
        for translations in TEXTS.values():
            self.assertNotIn("мезгил", translations["ky"].lower())

    def test_chart_stays_light_even_with_dark_matplotlib_style(self):
        series = ChartSeries([ChartBucket(date(2026, 10, 13), 200)], False)
        with plt.style.context("dark_background"):
            with chart_png(series, "99", "ky") as png:
                self.assertEqual(plt.gcf().axes[0].title.get_color(), "#1F1F1F")
                with Image.open(BytesIO(png)) as image:
                    self.assertEqual(image.convert("RGB").getpixel((0, 0)), (255, 255, 255))
        self.assertEqual(plt.get_fignums(), [])
