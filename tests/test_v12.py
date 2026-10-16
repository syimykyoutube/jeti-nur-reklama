"""v1.2 migration, required descriptions, rendering and unlimited reports."""

import asyncio
from dataclasses import replace
from html import escape, unescape
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from aiogram.methods import DeleteMessage, SendDocument, SendPhoto
from aiogram.types import Dice
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from bot.callbacks import AnalyticsPeriodCallback, ProductCallback, ReceiptCallback
from bot.db import Database, init_db
from bot.handlers.orders import OrderStates, order_card, receipt_caption
from bot.i18n import t
from bot.services.report_docx import create_report
from bot.services.stats import make_period
from test_stage1 import BotTestCase, NOW
from test_stage3 import DOCUMENT, PHOTOS


# Frozen pre-v1.2 schema: this fixture must not inherit future schema changes.
OLD_SCHEMA = """
CREATE TABLE users (tg_id INTEGER PRIMARY KEY, name TEXT, lang TEXT NOT NULL DEFAULT 'ky', created_at TEXT NOT NULL);
CREATE TABLE products (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, price REAL NOT NULL,
                      is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
CREATE TABLE orders (
    id INTEGER PRIMARY KEY AUTOINCREMENT, product_id INTEGER, product_name TEXT NOT NULL,
    price REAL NOT NULL, created_by INTEGER NOT NULL, created_by_name TEXT, created_at TEXT NOT NULL,
    receipt_file_id TEXT NOT NULL, receipt_file_unique_id TEXT, receipt_type TEXT NOT NULL,
    receipt_file_name TEXT, deleted_at TEXT, deleted_by INTEGER
);
CREATE INDEX idx_orders_created_at ON orders(created_at);
CREATE INDEX idx_orders_deleted_at ON orders(deleted_at);
"""


class MigrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_old_database_data_indexes_and_sequences_survive_repeated_init(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "old.db"
            with sqlite3.connect(path) as connection:
                connection.executescript(OLD_SCHEMA)
                connection.execute("INSERT INTO users VALUES (101, 'Сотрудник', 'ru', ?)", (NOW.isoformat(),))
                connection.execute("INSERT INTO products VALUES (42, 'Машина', 200, 0, ?)", (NOW.isoformat(),))
                for order_id, deleted_at, deleted_by in ((15, None, None), (19, NOW.isoformat(), 202)):
                    connection.execute("INSERT INTO orders VALUES (?, 42, 'Snapshot', 123.5, 101, 'Имя', ?, "
                                       "'file', 'unique', 'document', 'чек.pdf', ?, ?)",
                                       (order_id, NOW.isoformat(), deleted_at, deleted_by))
                before = {table: connection.execute(f"SELECT * FROM {table}").fetchall()
                          for table in ("users", "products", "orders", "sqlite_sequence")}
                old_columns = [row[1] for row in connection.execute("PRAGMA table_info(orders)")]
                objects = connection.execute("SELECT name, rootpage FROM sqlite_master ORDER BY name").fetchall()
            for _ in range(2):
                await init_db(path)
                with sqlite3.connect(path) as connection:
                    columns = connection.execute("PRAGMA table_info(orders)").fetchall()
                    self.assertEqual([row[1] for row in columns], old_columns + ["description"])
                    self.assertEqual(columns[-1][2], "TEXT")
                    self.assertEqual(connection.execute("SELECT description FROM orders").fetchall(), [(None,), (None,)])
                    self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone(), ("ok",))
                    for table in before:
                        projection = ", ".join(old_columns) if table == "orders" else "*"
                        self.assertEqual(connection.execute(f"SELECT {projection} FROM {table}").fetchall(), before[table])
                    self.assertEqual(connection.execute("SELECT name, rootpage FROM sqlite_master ORDER BY name").fetchall(), objects)
            order = await Database(path).get_active_order(15)
            self.assertIsNone(order.description)
            self.assertIsNone(await Database(path).get_active_order(19))
            new_id = await Database(path).add_order(
                product_id=42, product_name="New", price=200, created_by=101, created_by_name="Имя",
                receipt_file_id="new", receipt_file_unique_id="new-unique", receipt_type="photo",
                description="New description",
            )
            self.assertEqual(new_id, 20)

    async def test_concurrent_initialization_adds_column_once(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "old.db"
            with sqlite3.connect(path) as connection:
                connection.executescript(OLD_SCHEMA)
            await asyncio.gather(*(init_db(path) for _ in range(3)))
            with sqlite3.connect(path) as connection:
                self.assertEqual([row[1] for row in connection.execute("PRAGMA table_info(orders)")].count("description"), 1)


class DescriptionFlowTests(BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        await self.database.set_user_language(101, "Aibek", "ru")
        await self.database.set_user_language(202, "Aibek", "ky")
        self.product_id = await self.database.add_product("Машина", 200)

    def state(self, user_id=101):
        return self.dispatcher.fsm.get_context(bot=self.bot, chat_id=user_id, user_id=user_id)

    async def select(self, user_id=101):
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack(), user_id)

    async def rows(self):
        return (await self.database.get_orders_page()).orders

    def assert_not_deleted(self):
        self.assertFalse(any(isinstance(call, DeleteMessage) for call in self.session.calls))

    async def test_photo_and_document_captions_save_trimmed_description_in_both_languages(self):
        for user_id, lang, media in ((101, "ru", {"photo": PHOTOS}), (202, "ky", {"document": DOCUMENT})):
            await self.select(user_id)
            self.assertIn("описание", self.sent_messages()[-1].text)
            self.assertEqual(self.sent_messages()[-1].reply_markup.inline_keyboard[0][0].text, t("btn_back", lang))
            await self.send_message(None, user_id, caption="  Торговый центр <A> & Б  ", **media)
            order = (await self.rows())[0]
            self.assertEqual(order.description, "Торговый центр <A> & Б")
            self.assertEqual(order.created_by, user_id)
            self.assertEqual(order.receipt_type, "photo" if "photo" in media else "document")
            self.assertIsNone(await self.state(user_id).get_state())
            self.assertEqual(await self.state(user_id).get_data(), {})
            self.assertEqual(ProductCallback.unpack(self.sent_messages()[-1].reply_markup.inline_keyboard[0][0].callback_data).action, "select")
        self.assertEqual(len(await self.rows()), 2)
        self.assert_not_deleted()

    async def test_missing_empty_and_short_caption_wait_then_text_saves_in_both_languages(self):
        for user_id, lang in ((101, "ru"), (202, "ky")):
            for caption in (None, "  ", " ab "):
                before = len(await self.rows())
                await self.select(user_id)
                await self.send_message(None, user_id, document=DOCUMENT, caption=caption)
                self.assertEqual(len(await self.rows()), before)
                self.assertEqual(await self.state(user_id).get_state(), OrderStates.awaiting_description.state)
                data = await self.state(user_id).get_data()
                self.assertEqual((data["receipt_file_id"], data["receipt_file_unique_id"], data["receipt_type"], data["receipt_file_name"]),
                                 ("pdf", "pdf-unique", "document", "receipt.pdf"))
                self.assertEqual(self.sent_messages()[-1].text, t("ask_description", lang))
                self.assertEqual(self.sent_messages()[-1].reply_markup.inline_keyboard[0][0].callback_data, "order:cancel")
                await self.send_message("  Мен бул рекламаны соода борборунан алдым  ", user_id)
                self.assertEqual(len(await self.rows()), before + 1)
                self.assertEqual((await self.rows())[0].description, "Мен бул рекламаны соода борборунан алдым")
                self.assertIsNone(await self.state(user_id).get_state())
        self.assert_not_deleted()

    async def test_short_or_nontext_input_keeps_draft_and_does_not_save_or_delete_messages(self):
        for user_id, lang in ((101, "ru"), (202, "ky")):
            await self.select(user_id)
            await self.send_message(None, user_id, photo=PHOTOS)
            draft = await self.state(user_id).get_data()
            for text, content in (("ab", {}), ("  ", {}), (None, {"dice": Dice(emoji="🎲", value=2)}), ("/unknown", {})):
                await self.send_message(text, user_id, **content)
                self.assertEqual(self.sent_messages()[-1].text, t("description_wrong", lang))
                self.assertEqual(await self.state(user_id).get_data(), draft)
                self.assertEqual(await self.state(user_id).get_state(), OrderStates.awaiting_description.state)
                self.assertEqual(await self.rows(), [])
            await self.send_message("abc", user_id)
            self.assertEqual((await self.rows())[0].description, "abc")
            await self.database.soft_delete_order((await self.rows())[0].id, user_id)
        self.assert_not_deleted()

    async def test_long_caption_and_long_text_are_trimmed_to_700(self):
        for caption in (True, False):
            await self.select()
            await self.send_message(None, photo=PHOTOS, caption=" " + "Д" * 710 + " " if caption else None)
            if not caption:
                await self.send_message(" " + "Д" * 710 + " ")
            self.assertEqual((await self.rows())[0].description, "Д" * 700)

    async def test_new_receipt_replaces_all_metadata_before_description(self):
        await self.select()
        await self.send_message(None, document=DOCUMENT)
        await self.send_message(None, photo=PHOTOS, caption=" a ")
        draft = await self.state().get_data()
        self.assertEqual((draft["receipt_file_id"], draft["receipt_file_unique_id"], draft["receipt_type"], draft["receipt_file_name"]),
                         ("large", "large-unique", "photo", None))
        self.assertEqual(await self.rows(), [])
        await self.send_message("Новый источник")
        order = (await self.rows())[0]
        self.assertEqual((order.receipt_file_id, order.receipt_type, order.receipt_file_name), ("large", "photo", None))
        self.assert_not_deleted()

    async def test_new_receipt_with_valid_caption_saves_immediately(self):
        await self.select()
        await self.send_message(None, photo=PHOTOS)
        await self.send_message(None, document=DOCUMENT, caption="Источник рекламы")
        order = (await self.rows())[0]
        self.assertEqual((order.description, order.receipt_file_id, order.receipt_type, order.receipt_file_name),
                         ("Источник рекламы", "pdf", "document", "receipt.pdf"))
        self.assertEqual(len(await self.rows()), 1)
        self.assertIsNone(await self.state().get_state())

    async def test_cancel_and_back_at_both_steps_in_both_languages_discard_draft(self):
        for user_id, lang in ((101, "ru"), (202, "ky")):
            for with_receipt in (False, True):
                for action in ("/cancel", t("btn_back", lang), "order:cancel"):
                    await self.select(user_id)
                    if with_receipt:
                        await self.send_message(None, user_id, document=DOCUMENT)
                    if action == "order:cancel":
                        await self.click(action, user_id)
                    else:
                        await self.send_message(action, user_id)
                    self.assertIsNone(await self.state(user_id).get_state())
                    self.assertEqual(await self.state(user_id).get_data(), {})
                    self.assertEqual(await self.rows(), [])
                    self.assertEqual(ProductCallback.unpack(self.sent_messages()[-1].reply_markup.inline_keyboard[0][0].callback_data).action, "select")
                    await self.send_message("Late description", user_id)
                    self.assertEqual(await self.rows(), [])

    async def test_start_keeps_both_steps_and_supports_bot_mentions(self):
        for user_id in (101, 202):
            await self.select(user_id)
            for command in ("/start", "/start@test_bot payload"):
                before = await self.state(user_id).get_data()
                await self.send_message(command, user_id)
                self.assertEqual(await self.state(user_id).get_data(), before)
                self.assertEqual(await self.state(user_id).get_state(), OrderStates.awaiting_receipt.state)
            await self.send_message(None, user_id, document=DOCUMENT)
            before = await self.state(user_id).get_data()
            await self.send_message("/start", user_id)
            self.assertEqual(await self.state(user_id).get_data(), before)
            self.assertEqual(await self.state(user_id).get_state(), OrderStates.awaiting_description.state)
            await self.send_message("Source after /start", user_id)
            self.assertEqual((await self.rows())[0].description, "Source after /start")
        self.assert_not_deleted()

    async def test_snapshot_and_utc_timestamp_are_kept_until_description_arrives(self):
        await self.select()
        await self.send_message(None, photo=PHOTOS)
        await self.database.update_product_name(self.product_id, "Changed")
        await self.database.update_product_price(self.product_id, 999)
        await self.database.soft_delete_product(self.product_id)
        with patch("bot.db.datetime") as clock:
            clock.now.return_value = NOW
            await self.send_message("Original quoted product")
        order = (await self.rows())[0]
        self.assertEqual((order.product_name, order.price, order.created_at), ("Машина", 200, NOW.isoformat(timespec="microseconds")))

    async def test_concurrent_duplicate_description_creates_one_order(self):
        await self.select()
        await self.send_message(None, photo=PHOTOS)
        await asyncio.gather(self.send_message("First description"), self.send_message("Duplicate description"))
        self.assertEqual(len(await self.rows()), 1)

    async def test_database_refuses_new_order_without_valid_description(self):
        for description in (None, "", "   ", " x ", "ab"):
            with self.subTest(description=description), self.assertRaises(ValueError):
                await self.database.add_order(
                    product_id=self.product_id, product_name="Машина", price=200,
                    created_by=101, created_by_name="Aibek", receipt_file_id="file",
                    receipt_file_unique_id="unique", receipt_type="photo", description=description,
                )
        self.assertEqual(await self.rows(), [])


class DescriptionDisplayTests(BotTestCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        for user_id, lang in ((101, "ru"), (202, "ky")):
            await self.database.set_user_language(user_id, "Aibek", lang)
        self.product_id = await self.database.add_product("Видео <&>", 1000)

    async def order(self, description="Описание <b> & источник", kind="photo"):
        with patch("bot.db.datetime") as clock:
            clock.now.return_value = NOW
            order_id = await self.database.add_order(
                product_id=self.product_id, product_name="Видео <&>", price=1000,
                created_by=101, created_by_name="Автор <i>&", receipt_file_id="receipt",
                receipt_file_unique_id="unique", receipt_type=kind, receipt_file_name="receipt.pdf",
                description=description or "Temporary valid description",
            )
        if description is None:
            with sqlite3.connect(self.path) as connection:
                connection.execute("UPDATE orders SET description = NULL WHERE id = ?", (order_id,))
        return await self.database.get_active_order(order_id)

    async def test_analytics_and_deleteorder_use_same_escaped_card_in_both_languages(self):
        order = await self.order()
        for user_id, lang in ((101, "ru"), (202, "ky")):
            expected = ("<b>1. Видео &lt;&amp;&gt; — 1 000 сом</b>\nОписание &lt;b&gt; &amp; источник\n"
                        "<i>12:00 · 13.10.2026 · Автор &lt;i&gt;&amp;</i>")
            for action in ("delete", "list"):
                self.session.calls.clear()
                if action == "delete":
                    await self.send_message("/deleteorder", user_id)
                else:
                    await self.click(AnalyticsPeriodCallback(mode="orders", period="all").pack(), user_id)
                self.assertEqual(self.sent_messages()[0].text, expected)
                button = self.sent_messages()[0].reply_markup.inline_keyboard[0][0]
                self.assertEqual(button.text, t("delete_order_button" if action == "delete" else "btn_receipt", lang))
                self.assertEqual(len(self.sent_messages()), 1)
        self.assertEqual(self.bot.default.parse_mode, "HTML")

    async def test_old_null_description_has_no_empty_line(self):
        order = await self.order(None)
        for lang in ("ru", "ky"):
            card = order_card(order, 3, self.config, lang)
            self.assertEqual(card.count("\n"), 1)
            self.assertIn("<b>3.", card)
            self.assertNotIn("None", card)
            caption = receipt_caption(order, self.config, lang)
            self.assertEqual(caption.count("\n"), 1)
        await self.send_message("/deleteorder")
        self.assertEqual(self.sent_messages()[-1].text.count("\n"), 1)
        await self.click(AnalyticsPeriodCallback(mode="orders", period="all").pack())
        self.assertEqual(self.sent_messages()[-1].text.count("\n"), 1)

    async def test_card_limits_description_before_html_escape_and_adds_ellipsis(self):
        order = await self.order("<" * 700)
        line = order_card(order, 11, self.config, "ru").splitlines()[1]
        self.assertEqual(unescape(line), "<" * 299 + "…")
        self.assertNotIn("<", line)
        for length in (3, 299, 300):
            card = order_card(replace(order, description="a" * length), 1, self.config, "ky")
            self.assertEqual(card.splitlines()[1], "a" * length)

    async def test_receipt_photo_and_document_get_full_caption_in_both_languages(self):
        for kind in ("photo", "document"):
            order = await self.order(kind=kind)
            for user_id, lang in ((101, "ru"), (202, "ky")):
                await self.click(ReceiptCallback(order_id=order.id).pack(), user_id)
                sent = self.session.calls[-1]
                self.assertIsInstance(sent, SendPhoto if kind == "photo" else SendDocument)
                self.assertEqual(sent.caption, receipt_caption(order, self.config, lang))
                self.assertIn(escape(order.description), sent.caption)
                self.assertIn("<b>Видео &lt;&amp;&gt; — 1 000 сом</b>", sent.caption)
                self.assertIn("12:00 · 13.10.2026", sent.caption)
                self.assertIn(escape(order.created_by_name), sent.caption)

    async def test_caption_truncation_handles_emoji_and_html_without_cutting_entities(self):
        order = await self.order("😀" * 695 + "<&>")
        for kind in ("photo", "document"):
            for user_id, lang in ((101, "ru"), (202, "ky")):
                with sqlite3.connect(self.path) as connection:
                    connection.execute("UPDATE orders SET receipt_type = ? WHERE id = ?", (kind, order.id))
                await self.click(ReceiptCallback(order_id=order.id).pack(), user_id)
                caption = self.session.calls[-1].caption
                self.assertLessEqual(len(unescape(caption).encode("utf-16-le")) // 2, 1024)
                self.assertTrue(caption.splitlines()[1].endswith("…"))
                self.assertTrue(caption.endswith("</i>"))
                self.assertIn(escape(order.created_by_name), caption)
                self.assertNotIn("<&", caption)

    async def test_ranking_includes_every_product_sorts_and_splits_whole_html_lines(self):
        with sqlite3.connect(self.path) as connection:
            connection.executemany(
                "INSERT INTO orders (product_id, product_name, price, created_by, created_by_name, created_at, "
                "receipt_file_id, receipt_type, description) VALUES (?, ?, ?, 101, 'Author', ?, 'receipt', 'photo', 'Source')",
                [(self.product_id, f"{index:03d} " + "<&" * 25, index + 1, NOW.isoformat()) for index in range(100)],
            )
        period = make_period("all", self.config.timezone)
        top = await self.database.get_top_products(period)
        self.assertEqual(len(top), 100)
        self.assertEqual([row.total for row in top], list(range(100, 0, -1)))
        for user_id, lang in ((101, "ru"), (202, "ky")):
            self.session.calls.clear()
            await self.click(AnalyticsPeriodCallback(mode="top", period="all").pack(), user_id)
            messages = self.sent_messages()
            self.assertGreater(len(messages), 1)
            expected = [t("top_line", lang, rank=i, name=escape(item.name), count=1,
                          total=f"{item.total:g} сом") for i, item in enumerate(top, 1)]
            self.assertEqual("\n".join(m.text for m in messages).splitlines(), expected)
            for message in messages:
                self.assertLessEqual(len(message.text.encode("utf-16-le")) // 2, 4000)
                self.assertNotIn("<&", message.text)
            self.assertTrue(all(message.reply_markup is None for message in messages[:-1]))
            self.assertIsNotNone(messages[-1].reply_markup)

    async def test_word_has_description_all_products_all_orders_and_readable_layout(self):
        for index in range(15):
            with patch("bot.db.datetime") as clock:
                clock.now.return_value = NOW
                await self.database.add_order(
                    product_id=self.product_id, product_name=f"Product {index}", price=index + 1,
                    created_by=101, created_by_name="Автор & <>", receipt_file_id="PRIVATE_RECEIPT",
                    receipt_file_unique_id="PRIVATE_UNIQUE", receipt_type="photo",
                    description="Источник <&> " + "Длинное описание " * 35,
                )
        await self.order(None)
        deleted = await self.order("Deleted description")
        await self.database.soft_delete_order(deleted.id, 101)
        for lang in ("ru", "ky"):
            path = Path(self.temp.name) / f"report_{lang}.docx"
            await create_report(self.database, path, make_period("all", self.config.timezone, NOW), self.config.timezone, lang, NOW)
            document = Document(path)
            top, orders = document.tables[1:]
            self.assertIn(t("report_top", lang), [p.text for p in document.paragraphs])
            self.assertNotIn("топ 5", "\n".join(p.text for p in document.paragraphs).lower())
            self.assertEqual(len(top.rows), 17)  # 15 products plus one legacy product and header.
            self.assertEqual(len(orders.rows), 18)  # 16 orders, header and total.
            self.assertEqual([c.text for c in orders.rows[0].cells], [t(key, lang) for key in (
                "report_number", "report_date_column", "report_order_product", "report_amount", "report_description", "report_author")])
            self.assertEqual(orders.rows[1].cells[4].text, "—")
            self.assertIn("Источник <&>", orders.rows[2].cells[4].text)
            self.assertEqual(orders.rows[2].cells[5].text, "Автор & <>")
            self.assertEqual([round(column.width / Cm(1), 1) for column in orders.columns], [1, 3.2, 4, 2.2, 4, 2.6])
            section = document.sections[-1]
            self.assertLessEqual(sum(column.width for column in orders.columns), section.page_width - section.left_margin - section.right_margin)
            for table in (top, orders):
                self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn("w:tblHeader")))
                self.assertEqual(table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")), "1F4E79")
                self.assertEqual(table.rows[1].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")), "FFFFFF")
                self.assertEqual(table.rows[2].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")), "F5F7FA")
            self.assertTrue(orders.rows[-1].cells[0].text.startswith("Итого:" if lang == "ru" else "Жалпы:"))
            all_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            for absent in ("Deleted description", "PRIVATE_RECEIPT", "PRIVATE_UNIQUE"):
                self.assertNotIn(absent, all_text)
