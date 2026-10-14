"""Word contents, consistent snapshots, backup restoration, and Docker contract."""

import asyncio
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import sqlite3
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import aiosqlite
from aiogram.exceptions import TelegramBadRequest
from aiogram.methods import SendDocument
from aiogram.types import FSInputFile
from docx import Document

from bot.callbacks import AnalyticsKindCallback, AnalyticsPeriodCallback, PeriodChoiceCallback, ProductCallback
from bot.db import Database, init_db
from bot.i18n import t
from bot.services.backup import database_snapshot
from bot.services.charts import plt
from bot.services.report_docx import create_report, report_filename, xml_text
from bot.services.stats import make_period
from test_stage1 import BotTestCase, NOW


def document_text(document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


class ExportTests(BotTestCase):
    async def asyncSetUp(self) -> None:
        await super().asyncSetUp()
        await self.database.set_user_language(101, "Aibek", "ru")
        await self.database.set_user_language(202, "Aibek", "ky")
        self.product_id = await self.database.add_product("Машина", 200)

    async def order(self, name="Машина", price=200, when=NOW, author=101):
        order_id = await self.database.add_order(
            product_id=self.product_id, product_name=name, price=price, created_by=author,
            created_by_name=f"Author {author}", receipt_file_id="PRIVATE_RECEIPT_ID",
            receipt_file_unique_id="PRIVATE_UNIQUE_ID", receipt_type="document", receipt_file_name="PRIVATE_RECEIPT.pdf",
        )
        async with aiosqlite.connect(self.path) as connection:
            await connection.execute("UPDATE orders SET created_at = ? WHERE id = ?", (when.isoformat(), order_id))
            await connection.commit()
        return order_id

    async def export(self, period="all", user_id=101):
        with patch("bot.handlers.analytics.make_period", side_effect=lambda key, tz: make_period(key, tz, NOW)):
            await self.click(AnalyticsPeriodCallback(mode="word", period=period).pack(), user_id)
        filename, data = self.session.uploaded_documents[-1]
        return filename, data, Document(BytesIO(data))

    async def test_word_button_uses_selected_period_without_asking_again(self) -> None:
        await self.send_message(t("menu_analytics", "ru"))
        self.assertEqual(self.sent_messages()[-1].text, t("choose_period", "ru"))
        await self.click(PeriodChoiceCallback(period="d7").pack())
        keyboard = self.sent_messages()[-1].reply_markup.inline_keyboard
        button = keyboard[2][0]
        self.assertEqual(button.text, "📄 Word")
        payload = AnalyticsPeriodCallback.unpack(button.callback_data)
        self.assertEqual((payload.mode, payload.period), ("word", "d7"))
        self.assertEqual(self.session.uploaded_documents, [])
        await self.click(button.callback_data)
        self.assertEqual(len(self.session.uploaded_documents), 1)

    async def test_populated_report_contains_totals_tables_graph_and_no_receipts(self) -> None:
        for name, price in (("Машина", 200), ("Машина", 200), ("Үй", 300)):
            await self.order(name, price)
        await self.database.update_product_name(self.product_id, "New catalog name")
        filename, data, document = await self.export("7")
        text = document_text(document)
        self.assertTrue(filename.startswith("report_7_"))
        self.assertTrue(filename.endswith(".docx"))
        self.assertIn("Отчёт по заказам", text)
        self.assertIn("Количество заказов\n3", text)
        self.assertIn("Общая сумма\n700 сом", text)
        self.assertIn("Период: последние 7 дней", text)
        self.assertNotIn("Asia/Bishkek", text)
        self.assertEqual(len(document.tables), 3)
        totals, top, orders = document.tables
        self.assertEqual([cell.text for cell in top.rows[1].cells], ["1", "Машина", "2", "400 сом"])
        self.assertEqual(len(orders.rows), 5)
        self.assertEqual([cell.text for cell in orders.rows[1].cells], ["1", "12:00 · 13.10.2026", "Үй", "300 сом", "Author 101"])
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertNotIn("New catalog name", text)
        with ZipFile(BytesIO(data)) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith("word/media/")]), 1)
            xml = archive.read("word/document.xml").decode()
            for secret in ("PRIVATE_RECEIPT_ID", "PRIVATE_UNIQUE_ID", "PRIVATE_RECEIPT.pdf"):
                self.assertNotIn(secret, xml)
            self.assertIn("tblHeader", xml)
            self.assertIn("cantSplit", xml)
        self.assertEqual(plt.get_fignums(), [])

    async def test_empty_and_only_deleted_reports_still_sent(self) -> None:
        for user_id, lang in ((101, "ru"), (202, "ky")):
            _, _, document = await self.export("1", user_id)
            self.assertIn(t("no_orders", lang), document_text(document))
            self.assertEqual(len(document.tables), 0)
            self.assertEqual(len(document.inline_shapes), 0)
        deleted = await self.order()
        await self.database.soft_delete_order(deleted, 101)
        _, _, document = await self.export()
        self.assertIn(t("no_orders", "ru"), document_text(document))
        self.assertNotIn("Author 101", document_text(document))

    async def test_kyrgyz_document_fields_months_and_author(self) -> None:
        await self.order(author=202)
        _, _, document = await self.export(user_id=202)
        text = document_text(document)
        for expected in ("Заказдар боюнча отчёт", "Заказдын саны\n1", "Жалпы сумма\n200 сом", "Ким киргизди", "12:00 · 13.10.2026"):
            self.assertIn(expected, text)
        self.assertNotIn("Кто внёс", text)

    async def test_deleted_and_out_of_period_orders_are_absent(self) -> None:
        await self.order(name="Included", price=100)
        old = datetime(2026, 9, 1, tzinfo=timezone.utc)
        await self.order(name="Old", price=300, when=old)
        deleted = await self.order(name="Deleted", price=900)
        await self.database.soft_delete_order(deleted, 202)
        _, _, document = await self.export("1")
        text = document_text(document)
        self.assertIn("Общая сумма\n100 сом", text)
        self.assertIn("Included", text)
        self.assertNotIn("Deleted", text)
        self.assertNotIn("Old", text)

    async def test_all_order_rows_exported_across_batches(self) -> None:
        for index in range(105):
            await self.order(name=f"Product {index % 7}", price=index + 1)
        _, _, document = await self.export()
        self.assertEqual(len(document.tables[2].rows), 107)
        self.assertEqual(document.tables[2].rows[-2].cells[0].text, "105")
        self.assertEqual(len(document.tables[1].rows), 6)
        self.assertIn("Количество заказов\n105", document_text(document))

    async def test_report_sections_use_same_snapshot_during_concurrent_delete(self) -> None:
        order_id = await self.order()
        original = Database.get_order_summary

        async def mutate_after_summary(database, period):
            result = await original(database, period)
            if database.path != self.path:
                await self.database.soft_delete_order(order_id, 202)
            return result

        with patch.object(Database, "get_order_summary", new=mutate_after_summary):
            _, _, document = await self.export()
        self.assertIsNone(await self.database.get_active_order(order_id))
        self.assertIn("Количество заказов\n1", document_text(document))
        self.assertEqual(len(document.tables[2].rows), 3)
        self.assertEqual(len(document.inline_shapes), 1)

    async def test_backup_sends_copy_and_restore_has_all_records(self) -> None:
        order_id = await self.order()
        await self.database.soft_delete_order(order_id, 202)
        await self.send_message("/backup")
        delivery = next(call for call in self.session.calls if isinstance(call, SendDocument))
        self.assertIsInstance(delivery.document, FSInputFile)
        self.assertNotEqual(Path(delivery.document.path), self.path)
        self.assertFalse(Path(delivery.document.path).exists())
        filename, data = self.session.uploaded_documents[-1]
        self.assertTrue(filename.startswith("backup_"))
        restored = Path(self.temp.name) / "restored.db"
        restored.write_bytes(data)
        with sqlite3.connect(restored) as connection:
            self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM users").fetchone()[0], 2)
            self.assertEqual(connection.execute("SELECT receipt_file_id, deleted_by FROM orders").fetchone(), ("PRIVATE_RECEIPT_ID", 202))
        self.assertIsNone(await self.database.get_active_order(order_id))
        await init_db(restored)
        self.assertEqual((await Database(restored).list_active_products())[0].name, "Машина")

    async def test_backup_reads_committed_wal_data(self) -> None:
        # Keep this connection open so committed changes remain in the WAL sidecar.
        async with aiosqlite.connect(self.path) as writer:
            await writer.execute("PRAGMA journal_mode=WAL")
            await writer.execute("PRAGMA wal_autocheckpoint=0")
            order_id = await self.order(name="WAL row")
            async with database_snapshot(self.database) as snapshot:
                self.assertEqual((await snapshot.get_active_order(order_id)).product_name, "WAL row")
                snapshot_path = snapshot.path
            self.assertFalse(snapshot_path.exists())

    async def test_backup_preserves_receipt_waiting_state(self) -> None:
        await self.click(ProductCallback(action="select", product_id=self.product_id).pack())
        state = self.dispatcher.fsm.get_context(bot=self.bot, chat_id=101, user_id=101)
        before = await state.get_data()
        await self.send_message("/backup")
        self.assertEqual(await state.get_data(), before)
        self.assertIsNotNone(await state.get_state())

    async def test_whitelist_blocks_word_and_backup(self) -> None:
        await self.send_message("/backup", 999)
        await self.click(AnalyticsKindCallback(mode="word").pack(), 999)
        await self.click(AnalyticsPeriodCallback(mode="word", period="all").pack(), 999)
        self.assertEqual(self.session.calls, [])
        self.assertEqual(self.session.uploaded_documents, [])

    async def test_backup_failure_returns_localized_message(self) -> None:
        with patch.object(Database, "backup_to", side_effect=sqlite3.OperationalError("locked")), self.assertLogs(level="ERROR"):
            await self.send_message("/backup", 202)
        self.assertEqual(self.sent_messages()[-1].text, t("backup_failed", "ky"))
        self.assertEqual(self.session.uploaded_documents, [])

    async def test_send_failures_cleanup_temporary_files(self) -> None:
        paths = []
        original = self.session.make_request

        async def unavailable(bot, method, timeout=None):
            if isinstance(method, SendDocument):
                paths.append(Path(method.document.path))
                self.assertTrue(paths[-1].exists())
                raise TelegramBadRequest(method=method, message="send failed")
            return await original(bot, method, timeout)

        with patch.object(self.session, "make_request", side_effect=unavailable), self.assertLogs(level="ERROR"):
            await self.send_message("/backup")
            with patch("bot.handlers.analytics.make_period", side_effect=lambda key, tz: make_period(key, tz, NOW)):
                await self.click(AnalyticsPeriodCallback(mode="word", period="all").pack())
        self.assertEqual(len(paths), 2)
        self.assertTrue(all(not path.exists() for path in paths))
        self.assertTrue(self.path.exists())

    async def test_backup_refuses_working_file_and_existing_destination(self) -> None:
        with self.assertRaises(ValueError):
            await self.database.backup_to(self.path)
        destination = Path(self.temp.name) / "existing.db"
        destination.write_bytes(b"untouched")
        with self.assertRaises(FileExistsError):
            await self.database.backup_to(destination)
        self.assertEqual(destination.read_bytes(), b"untouched")

    async def test_backup_while_writes_remains_readable(self) -> None:
        destination = Path(self.temp.name) / "concurrent.db"

        async def write_orders():
            for _ in range(8):
                await self.order()

        await asyncio.gather(self.database.backup_to(destination), write_orders())
        with sqlite3.connect(destination) as connection:
            self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertLessEqual(connection.execute("SELECT COUNT(*) FROM orders").fetchone()[0], 8)

    async def test_report_cancellation_waits_for_writer_before_cleanup(self) -> None:
        output = Path(self.temp.name) / "cancelled.docx"
        started = asyncio.Event()
        release = asyncio.Event()

        async def slow_thread(*args):
            started.set()
            await release.wait()
            return args[0](*args[1:])

        with patch("bot.services.report_docx.asyncio.to_thread", side_effect=slow_thread):
            task = asyncio.create_task(create_report(self.database, output, make_period("all", self.config.timezone, NOW), self.config.timezone, "ru", NOW))
            await started.wait()
            task.cancel()
            release.set()
            with self.assertRaises(asyncio.CancelledError):
                await task
        self.assertTrue(output.exists())
        self.assertIn(t("no_orders", "ru"), document_text(Document(output)))


class ReportAndDockerTests(unittest.TestCase):
    def test_filename_and_xml_cleanup(self) -> None:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo("Asia/Bishkek")
        self.assertEqual(report_filename(make_period("7", tz, NOW), tz, NOW), "report_7_2026-10-13.docx")
        self.assertEqual(xml_text("A\x00B\nКыргызча"), "AB\nКыргызча")

    def test_dockerfile_and_secrets_exclusion(self) -> None:
        root = Path(__file__).resolve().parent.parent
        dockerfile = (root / "Dockerfile").read_text()
        self.assertIn("FROM python:3.11-slim", dockerfile)
        self.assertIn("DB_PATH=/data/bot.db", dockerfile)
        self.assertIn('VOLUME ["/data"]', dockerfile)
        self.assertIn('CMD ["python", "-m", "bot.main"]', dockerfile)
        self.assertIn("COPY bot/ ./bot/", dockerfile)
        self.assertNotIn("COPY .", dockerfile)
        ignores = (root / ".dockerignore").read_text().splitlines()
        for name in (".env", ".env.*", "data", ".venv", ".runtime", "*.db"):
            self.assertIn(name, ignores)
