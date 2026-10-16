"""Localized reports generated from a consistent database snapshot, without receipts."""

import asyncio
from contextlib import closing
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import sqlite3
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.run import Run
from docx.shared import Cm, Pt, RGBColor

from bot.db import Database, order_scope
from bot.i18n import MONTHS, format_datetime, format_money, normalize_language, t
from bot.services.backup import database_snapshot
from bot.services.charts import chart_png
from bot.services.stats import Period, Summary, TopProduct, chart_weekly, fill_chart_buckets, period_label


def xml_text(value: object) -> str:
    """Telegram names may contain characters forbidden in Word XML."""
    return "".join(char for char in str(value) if (
        char in "\t\n\r" or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF
    ))


def report_datetime(value: datetime, lang: str, tz: ZoneInfo) -> str:
    local = value.astimezone(tz)
    return t("report_datetime", lang, day=local.day, month=MONTHS[normalize_language(lang)][local.month - 1],
             year=local.year, time=f"{local:%H:%M}")


def report_filename(period: Period, tz: ZoneInfo, generated_at: datetime) -> str:
    return f"report_{period.key}_{generated_at.astimezone(tz):%Y-%m-%d}.docx"


TEXT_COLOR = "1F1F1F"
ACCENT = "1F4E79"


def style_run(run, *, size=11, bold=False, color=TEXT_COLOR):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key, value in (("ascii", "Calibri"), ("hAnsi", "Calibri"), ("eastAsia", "Arial"), ("cs", "Arial")):
        fonts.set(qn(f"w:{key}"), value)
    return run


def paragraph(document, text, *, style=None, center=False, size=11, bold=False, color=TEXT_COLOR):
    result = document.add_paragraph(style=style)
    if center:
        result.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(result.add_run(xml_text(text)), size=size, bold=bold, color=color)
    return result


def configure_document(document, lang: str, generated_at: datetime) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    section.footer_distance = Cm(1)
    background = OxmlElement("w:background")
    background.set(qn("w:color"), "FFFFFF")
    document._element.insert(0, background)
    display = OxmlElement("w:displayBackgroundShape")
    display.set(qn("w:val"), "true")
    document.settings.element.append(display)
    for name, size in (("Normal", 11), ("Title", 22), ("Heading 1", 14), ("Caption", 10)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
        if name == "Heading 1":
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "ru-RU" if lang == "ru" else "ky-KG")
    document.styles["Normal"].element.get_or_add_rPr().append(language)
    # The alternate family is honored by applications that do not have Calibri.
    for relation in document.part.rels.values():
        if relation.reltype.endswith("/fontTable"):
            from lxml import etree
            root = etree.fromstring(relation.target_part.blob)
            fonts = root.findall(qn("w:font"))
            font = next((item for item in fonts if item.get(qn("w:name")) == "Calibri"), None)
            if font is None:
                font = OxmlElement("w:font")
                font.set(qn("w:name"), "Calibri")
                root.append(font)
            alternate = OxmlElement("w:altName")
            alternate.set(qn("w:val"), "Arial")
            font.append(alternate)
            relation.target_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field_run = style_run(footer.add_run(), size=10)
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    field_run._r.append(field)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    style_run(footer.add_run(), size=10)._r.append(instruction)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    style_run(footer.add_run(), size=10)._r.append(separator)
    style_run(footer.add_run("1"), size=10)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    style_run(footer.add_run(), size=10)._r.append(end)
    document.core_properties.title = t("report_title", lang)
    document.core_properties.author = document.core_properties.last_modified_by = "Jeti Nur"
    document.core_properties.created = generated_at.astimezone(timezone.utc).replace(tzinfo=None)
    document.core_properties.modified = document.core_properties.created


def shade(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def table_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(sum(widths) / 2.54 * 1440)))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D0D7DE")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "80"), ("bottom", "80"), ("left", "100"), ("right", "100")):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)


def make_table(document, headers: list[str], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table_layout(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    fill_row(table.rows[0], headers, widths, header=True)
    return table


def fill_row(row, values: list[str], widths: list[float], *, header=False, alignments=None, fill="FFFFFF", bold=False):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, ACCENT if header else fill)
        cell_paragraph = cell.paragraphs[0]
        cell_paragraph.paragraph_format.space_after = Pt(0)
        cell_paragraph.paragraph_format.line_spacing = 1.05
        cell_paragraph.paragraph_format.keep_with_next = header
        cell_paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if header else
                                   (alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT))
        style_run(cell_paragraph.add_run(xml_text(value)), size=11, bold=header or bold,
                  color="FFFFFF" if header else TEXT_COLOR)


def write_report(snapshot_path: Path, output: Path, period: Period, tz: ZoneInfo, lang: str,
                 generated_at: datetime, summary: Summary, top: list[TopProduct], png: bytes | None,
                 bot_username: str | None = None) -> None:
    document = Document()
    configure_document(document, lang, generated_at)
    title = paragraph(document, t("report_title", lang), style="Title", center=True, size=22, bold=True, color=ACCENT)
    title.paragraph_format.keep_with_next = True
    paragraph(document, t("report_subtitle", lang, period=period_label(period.key, lang),
                          date=report_datetime(generated_at, lang, tz)), center=True)
    if not summary.count:
        paragraph(document, t("no_orders", lang))
        document.save(output)
        return

    totals = document.add_table(rows=1, cols=2)
    table_layout(totals, [8.5, 8.5])
    for cell, key, value in zip(totals.rows[0].cells, ("report_count_label", "report_total_label"),
                               (str(summary.count), format_money(summary.total, lang))):
        cell.width = Cm(8.5)
        shade(cell, "EAF1FB")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        entry = cell.paragraphs[0]
        style_run(entry.add_run(t(key, lang) + "\n"))
        style_run(entry.add_run(value), size=18, bold=True, color=ACCENT)

    paragraph(document, t("report_top", lang), style="Heading 1", size=14, bold=True)
    widths = [1.2, 9.8, 2.5, 3.5]
    top_table = make_table(document, [t(key, lang) for key in (
        "report_number", "report_product", "report_quantity", "report_amount")], widths)
    for index, item in enumerate(top, 1):
        fill_row(top_table.add_row(), [str(index), item.name, str(item.count), format_money(item.total, lang)], widths,
                 alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
                 fill="FFFFFF" if index % 2 else "F5F7FA")

    paragraph(document, t("report_chart", lang), style="Heading 1", size=14, bold=True)
    if png is not None:
        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.keep_with_next = True
        image_run = style_run(picture_paragraph.add_run())
        with BytesIO(png) as image:
            image_run.add_picture(image, width=Cm(16))
        document.inline_shapes[-1]._inline.docPr.set("descr", t("report_chart", lang))
        paragraph(document, t("report_chart", lang), style="Caption", center=True, size=10)

    widths = [1, 3, 5, 2.4, 6, 3.5, 1.8]
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    paragraph(document, t("report_orders", lang), style="Heading 1", size=14, bold=True)
    table = make_table(document, [t(key, lang) for key in (
        "report_number", "report_date_column", "report_order_product", "report_amount",
        "report_description", "report_author", "report_receipt")], widths)
    scope, values = order_scope(period.start, period.end)
    # Only display columns, read in batches from the same consistent snapshot.
    with closing(sqlite3.connect(snapshot_path.resolve().as_uri() + "?mode=ro", uri=True)) as connection:
        cursor = connection.execute(
            f"SELECT product_name, price, created_by_name, created_by, created_at, description, id, "
            f"receipt_file_id, receipt_type FROM orders WHERE {scope} "
            "ORDER BY created_at DESC, id DESC", values)
        index = 0
        while rows := cursor.fetchmany(100):
            for name, price, author, author_id, created_at, description, order_id, receipt_id, receipt_type in rows:
                index += 1
                row = table.add_row()
                has_link = bool(bot_username and receipt_id and receipt_type in {"photo", "document"})
                fill_row(row, [str(index), format_datetime(datetime.fromisoformat(created_at), lang, tz),
                                         name, format_money(price, lang),
                                         (description or "").strip() or t("report_description_empty", lang),
                                         author or str(author_id), "" if has_link else t("report_description_empty", lang)], widths,
                         alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                                     WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
                                     WD_ALIGN_PARAGRAPH.CENTER],
                         fill="FFFFFF" if index % 2 else "F5F7FA")
                if has_link:
                    link_paragraph = row.cells[-1].paragraphs[0]
                    relationship = link_paragraph.part.relate_to(
                        f"https://t.me/{bot_username}?start=r{order_id}", RT.HYPERLINK, is_external=True)
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), relationship)
                    run = style_run(Run(OxmlElement("w:r"), link_paragraph), color="0563C1")
                    run.font.underline = True
                    run.text = t("report_receipt", lang)
                    hyperlink.append(run._r)
                    link_paragraph._p.append(hyperlink)
    last_row = table.add_row()
    merged = last_row.cells[0].merge(last_row.cells[-1])
    shade(merged, "EAF1FB")
    merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(merged.paragraphs[0].add_run(t("report_grand_total", lang, total=format_money(summary.total, lang))), bold=True)
    last_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    document.save(output)


async def create_report(
    database: Database, output: Path, period: Period, tz: ZoneInfo, lang: str,
    generated_at: datetime | None = None,
    *, bot_username: str | None = None,
) -> None:
    generated_at = generated_at if generated_at is not None else datetime.now(timezone.utc)
    lang = normalize_language(lang)
    async with database_snapshot(database) as snapshot:
        summary = await snapshot.get_order_summary(period)
        top = await snapshot.get_top_products(period)
        png = None
        if summary.count:
            weekly = chart_weekly(period, summary.first_created_at, tz)
            rows = await snapshot.get_chart_buckets(period, tz, weekly)
            series = fill_chart_buckets(rows, period, tz, weekly)
            with chart_png(series, period.key, lang) as data:
                png = data
        # All matplotlib work stays on the event-loop thread. Word assembly is offloaded.
        task = asyncio.create_task(asyncio.to_thread(
            write_report, snapshot.path, output, period, tz, lang, generated_at, summary, top, png, bot_username,
        ))
        try:
            await asyncio.shield(task)
        except asyncio.CancelledError:
            # Let the writer finish before its temporary snapshot directory is removed.
            await task
            raise
